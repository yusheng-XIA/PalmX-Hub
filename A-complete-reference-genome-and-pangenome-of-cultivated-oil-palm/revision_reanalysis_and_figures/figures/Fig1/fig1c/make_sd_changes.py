#!/usr/bin/env python3
"""Build the Fig. 1c Source Data cell-change list (file, sheet, cell, old_value, new_value) from the split
workbook. Output: sd_changes_fig1c.tsv (input for deliver_build/apply_cell_changes.py)."""
import csv
from pathlib import Path
import openpyxl

S = Path(__file__).resolve().parents[2]
SPLIT = S / "deliver/Source_Data_split/Source_Data_Fig1.xlsx"
OUT = Path(__file__).with_name("sd_changes_fig1c.tsv")
FILE = "Source_Data_Fig1.xlsx"

# exact-cell replacements (sheet, cell, new); old value is read from the workbook
EXPLICIT = [
    ("Fig1c_TE_density", "D2", "TE coverage (fraction)"),
    ("Fig1c_gene_density", "F2", "FL-specific gene number"),
]
# substring replacements inside Fig1c_* sheets (applied to every string cell that contains them)
SUBS = [
    ("FL Africa hap2", "FL-Hap2"),
    ("Annotation-missing genes have no EG11 orthogroup member",
     "FL-specific genes are FL genes lacking an EG11 orthologue in the annotation (no EG11 orthogroup member"),
    ("and no qualifying DIAMOND protein hit;", "and no qualifying DIAMOND protein hit);"),
    ("The annotation-missing count is a subset", "The FL-specific count is a subset"),
]
CONTENTS = [("telomeres", "chromosome ends")]

wb = openpyxl.load_workbook(SPLIT)
rows = []
for sh, cell, new in EXPLICIT:
    rows.append((FILE, sh, cell, wb[sh][cell].value, new))
done = {(r[1], r[2]) for r in rows}
for sh in wb.sheetnames:
    if not sh.startswith("Fig1c"):
        continue
    for row in wb[sh].iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or (sh, c.coordinate) in done:
                continue
            nv = v
            for a, b in SUBS:
                nv = nv.replace(a, b)
            if nv != v:
                rows.append((FILE, sh, c.coordinate, v, nv))
cont = wb["Contents"]
for old, new in CONTENTS:
    hits = [i for i in range(1, cont.max_row + 1) if cont.cell(i, 4).value == old]
    assert len(hits) == 1, (old, hits)
    rows.append((FILE, "Contents", f"D{hits[0]}", old, new))
with open(OUT, "w", newline="") as f:
    w = csv.writer(f, delimiter="\t", lineterminator="\n")
    w.writerow(["file", "sheet", "cell", "old_value", "new_value"])
    w.writerows(rows)
print(OUT, len(rows), "rows")
