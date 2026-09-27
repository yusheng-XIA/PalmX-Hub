#!/usr/bin/env python3
"""Source Data Fig. 1a: Descals tiles with recomputed oil-palm cover (duplicate tile row removed) and the plotted
0.05-degree cover grid, both re-aggregated from Zenodo 4473715 (author's original plotting table not found).
usage: fix_sd_1a.py Source_Data.xlsx fix/fig1a"""
import csv, re, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
p, fx = sys.argv[1], Path(sys.argv[2])
wb = openpyxl.load_workbook(p); cont = wb["Contents"]
def val(v):
    if v == "": return None
    if re.fullmatch(r"-?\d+", v): return int(v)
    try: return float(v)
    except ValueError: return v
def crow(n):
    for r in range(1, cont.max_row + 1):
        if cont.cell(r, 1).value == n: return r
def put(name, tsv, desc, replace=None, after=None):
    t = replace or name
    if t in wb.sheetnames:
        i = wb.sheetnames.index(t); del wb[t]; r = crow(t)
    else:
        i = wb.sheetnames.index(after) + 1; r = crow(after) + 1; cont.insert_rows(r)
    ws = wb.create_sheet(name, i)
    for k, row in enumerate(csv.reader(open(tsv), delimiter="\t")):
        ws.append([val(v) for v in row])
        if k == 0:
            for c in ws[1]: c.font = Font(bold=True)
    for j, v in enumerate([name, "Figure 1", "a", desc], 1): cont.cell(r, j).value = v
    print(name, ws.max_row)
put("Fig1a_detection_grid", fx / "Fig1a_cover_tiles_dedup.tsv",
    "633 Descals et al. 2019 map tiles with oil palm (duplicate tile row removed) and oil-palm, industrial and "
    "smallholder cover recomputed from the 10-m classification (Zenodo 4473715)", replace="Fig1a_detection_grid")
put("Fig1a_cover_grid_005deg", fx / "Fig1a_cover_grid.tsv",
    "Plotted grid: oil-palm cover per 0.05-degree cell (cells with >=1% cover are coloured on a log scale)",
    after="Fig1a_detection_grid")
wb.save(p); print("saved", len(wb.sheetnames))
