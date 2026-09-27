import openpyxl, math
A='fix/fig1c_author/Source_Data_Fig1c.xlsx'; B='deliver/Source_Data_split/Source_Data_Fig1.xlsx'
wa=openpyxl.load_workbook(A,read_only=True); wb=openpyxl.load_workbook(B,read_only=True)
pairs=[(f'Source Data Fig1c ({i})',n) for i,n in enumerate(['Fig1c_chr_lengths','Fig1c_gene_density','Fig1c_TE_density','Fig1c_EG11_gaps','Fig1c_segmental_dups','Fig1c_tandem_repeats','Fig1c_telomeres','Fig1c_synteny'],1)]
def norm(v):
    if isinstance(v,str): return v.replace('FL Africa hap2','FL-Hap2')
    return v
for a,b in pairs:
    ra=list(wa[a].iter_rows(values_only=True)); rb=list(wb[b].iter_rows(values_only=True))
    da=[r for r in ra[2:] if any(x is not None for x in r)][:-1]
    db=[r for r in rb[2:] if any(x is not None for x in r)][:-1]
    nd=0; ex=[]
    for x,y in zip(da,db):
        x=tuple(norm(v) for v in x); y=tuple(y)
        same=all((isinstance(p,(int,float)) and isinstance(q,(int,float)) and math.isclose(p,q,rel_tol=1e-9,abs_tol=1e-12)) or p==q for p,q in zip(x,y))
        if not same:
            nd+=1
            if len(ex)<3: ex.append((x,y))
    print(f'{a} vs {b}: rows {len(da)} / {len(db)}; diff rows {nd}; header A={ra[1]} B={rb[1]}')
    for e in ex: print('   ',e)
    # Sample value set
    if ra[1][0]=='Sample':
        from collections import Counter; print('   samples',Counter(r[0] for r in da))
