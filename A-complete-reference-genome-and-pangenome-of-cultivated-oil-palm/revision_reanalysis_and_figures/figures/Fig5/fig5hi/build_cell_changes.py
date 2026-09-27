"""List Source Data (Fig. 5) and Supplementary Data 24 cells to change for the W = 0 capture recompute.
Reads deliver_final workbooks read-only; writes table_cell_changes_5hi.tsv here."""
import csv, re
from pathlib import Path
import openpyxl
H = Path(__file__).resolve().parent; S = H.parents[1]
SUB = S / "deliver_final/OilPalm_Nature_submission_20260924"
SD24 = SUB / "05_Supplementary_Data/Supplementary_Data_24.xlsx"
SDF5 = SUB / "06_Source_Data/Source_Data_Fig5.xlsx"
def disp(d):
    m = re.fullmatch(r"(dura|pisifera|nrly|bk)_hap([12])", d)
    return f"{dict(dura='TK', pisifera='NS', nrly='Nigerian', bk='TN')[m.group(1)]}-Hap{m.group(2)}" if m else d
loci = {r["SV"]: r for r in csv.DictReader(open(H / "capture_current_path.tsv"), delimiter="\t")}
per = {r["Chrom"]: r for r in csv.DictReader(open(H / "per_chrom_capture.tsv"), delimiter="\t")}
rows = []
def add(f, sh, cell, old, new, need, note=""):
    if str(old) != str(new): rows.append((f, sh, cell, old, new, need, note))
# check donor display naming against the Source Data window path
wb = openpyxl.load_workbook(SDF5)
wp = wb["Fig5h_window_path"]; raw = list(csv.DictReader(open(H.parent / "misc9/data/ideal_loadonly_path.tsv"), delimiter="\t"))
mism = sum(1 for r, x in zip(raw, wp.iter_rows(min_row=2, values_only=True)) if disp(r["Donor_ID"]) != x[4])
print("window-path donor naming mismatches:", mism)
F = "Source_Data_Fig5.xlsx"
c = wb["Contents"]
for cell, new in [("D16", "Per-chromosome summary of the donor path (W = 0, P = 15), with favourable-locus capture"),
                  ("D17", "Window-level donor path (W = 0, P = 15; 3,486 windows)"),
                  ("D18", "284 SV-GWAS favourable loci and their capture by the W = 0, P = 15 path"),
                  ("D19", "chr01B windows selected by the W = 0, P = 15 path")]:
    add(F, "Contents", cell, c[cell].value, new, "required" if cell in ("D16", "D18") else "recommended")
s = wb["Fig5h_chromosome_summary"]
add(F, s.title, "K1", s["K1"].value, "Fav_Captured", "required"); add(F, s.title, "L1", s["L1"].value, "Fav_Total", "required")
add(F, s.title, "N1", s["N1"].value, "Fav_No_Donor_Call", "optional", "new column; counted as not captured")
for r in range(2, s.max_row + 1):
    ch = s[f"A{r}"].value; p = per[ch]
    assert int(p["Fav_Total"]) == s[f"L{r}"].value, ch
    add(F, s.title, f"K{r}", s[f"K{r}"].value, int(p["Captured_W0"]), "required")
    if int(p["Fav_Total"]):
        add(F, s.title, f"M{r}", s[f"M{r}"].value, round(100 * int(p["Captured_W0"]) / int(p["Fav_Total"]), 2), "required")
    add(F, s.title, f"N{r}", None, int(p["Unknown_W0_no_donor_call"]), "optional")
def loci_sheet(f, sh, r0, colE, colF, colG):
    # r0 = header row
    add(f, sh.title, f"{colE}{r0}", sh[f"{colE}{r0}"].value, "Selected_Donor", "required")
    add(f, sh.title, f"{colG}{r0}", sh[f"{colG}{r0}"].value, "Donor_Genotype_Call", "optional", "new column")
    n = 0; r = r0 + 1
    while sh[f"A{r}"].value and str(sh[f"A{r}"].value).startswith("SV_"):
        L = loci[sh[f"A{r}"].value]; n += 1
        assert int(sh[f"C{r}"].value) == int(L["Pos"]) and sh[f"D{r}"].value == L["Target"]
        add(f, sh.title, f"{colE}{r}", sh[f"{colE}{r}"].value, disp(L["Current_W0_Selected_Donor"]), "required")
        add(f, sh.title, f"{colF}{r}", sh[f"{colF}{r}"].value, int(L["Captured_W0"]), "required")
        add(f, sh.title, f"{colG}{r}", None, "No_call" if L["Capture_Status_W0"].startswith("Unknown") else "Called", "optional")
        r += 1
    return n, r
n, _ = loci_sheet(F, wb["Fig5h_favourable_loci"], 1, "E", "F", "G"); print("SourceData loci rows", n)
# ---- SD24
G = "Supplementary_Data_24.xlsx"
w = openpyxl.load_workbook(SD24); t = w.active
fixed = [
    ("A1", "Supplementary Data 24. Donor path (W = 0, P = 15), candidate deleterious load and favourable-locus capture", "required", "also titles.tsv row 24 and any SI list of Supplementary Data titles"),
    ("A2", "Donor-path result (W = 0, P = 15)", "required", ""),
    ("A4", "African35 donor path (W = 0, P = 15)", "required", ""),
    ("C4", "35 haplotypes (genotype calls for 29)", "required", "favourable loci now evaluated on the 35-haplotype path"),
    ("F4", 0, "required", "objective I52 = 28,710 = 23,520 + 15 x 346, no GWAS term"),
    ("A7", "Inputs and exact results (W = 0, P = 15)", "required", ""),
    ("B10", 35, "required", ""),
    ("D10", "Genotype calls from the graph-pangenome bubble VCF for 29 haplotypes; TK, NS and Nigerian haplotypes not represented", "required", ""),
    ("B17", 0, "required", ""),
    ("D17", "GWAS alleles not included in the path objective", "required", ""),
    ("B29", 168, "required", ""),
    ("D29", "Sum of Captured (selected donor carries the favourable allele)", "required", ""),
    ("A32", "Favourable loci without a selected-donor genotype call", "optional", "new row (row 32 is empty)"),
    ("B32", 19, "optional", ""), ("C32", "loci", "optional", ""), ("D32", "Counted as not captured", "optional", ""),
    ("B31", 59.15, "required", ""),
    ("D31", "168 / 284 × 100", "required", ""),
    ("A34", "Chromosome-level summary (W = 0, P = 15)", "required", ""),
    ("K35", "Fav_Captured", "required", ""), ("L35", "Fav_Total", "required", ""),
    ("A54", "Window-level donor path, W = 0, P = 15 (3,486 rows)", "required", "3,486/3,486 windows = results-8.9 african35/ideal_loadonly_path.tsv"),
    ("A3544", "Locus-level favourable-allele capture on the W = 0, P = 15 path (284 rows)", "required", ""),
    ("A3830", "Note: Captured = 1 when the donor selected for the 500-kb window containing the locus carries the favourable (target) allele; otherwise 0. Loci whose selected donor has no genotype call (TK, NS and Nigerian haplotypes, which are not represented in the graph-pangenome bubble VCF, or a missing call) are counted as 0. Donor names follow the window-level path above.", "required", ""),
]
for cell, new, need, note in fixed: add(G, t.title, cell, t[cell].value, new, need, note)
for r in range(36, 52):
    ch = t[f"A{r}"].value; p = per[ch]
    assert int(p["Fav_Total"]) == t[f"L{r}"].value
    add(G, t.title, f"K{r}", t[f"K{r}"].value, int(p["Captured_W0"]), "required")
    if int(p["Fav_Total"]): add(G, t.title, f"M{r}", t[f"M{r}"].value, round(100 * int(p["Captured_W0"]) / int(p["Fav_Total"]), 2), "required")
add(G, t.title, "K52", t["K52"].value, 168, "required"); add(G, t.title, "M52", t["M52"].value, 59.15, "required")
n, rend = loci_sheet(G, t, 3545, "E", "F", "G"); print("SD24 loci rows", n, "note row", rend)
with open(H / "table_cell_changes_5hi.tsv", "w", newline="") as fh:
    wr = csv.writer(fh, delimiter="\t", lineterminator="\n")
    wr.writerow(["file", "sheet", "cell", "old_value", "new_value", "necessity", "note"]); wr.writerows(rows)
from collections import Counter
print(Counter((r[0], r[5]) for r in rows))
