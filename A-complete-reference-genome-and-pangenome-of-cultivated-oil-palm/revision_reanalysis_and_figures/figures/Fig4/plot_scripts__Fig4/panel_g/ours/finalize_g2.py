#!/usr/bin/env python3
"""G2 finalisation: ST19 rows, Source Data TSV and Fig. 4g bars from the final-annotation rerun.

usage: finalize_g2.py SUMMARY_DIR [--dry]
Inputs: <pair>_final.allele_summary.final_annot_20260923.tsv for dura, pisifera, nrly, MZ4.
Category mapping (unchanged from 06_build_g_allele_plot.py):
  Bi-allelic = Heterozygous; Haplotype-specific = Unique_inter; Same CDS = Homozygous; Unresolved = No_hit.
"""
import csv, shutil, sys
from pathlib import Path

import fitz
import openpyxl
from PIL import Image

S = Path(__file__).resolve().parents[1]           # scratchpad
SUM = Path(sys.argv[1]); DRY = "--dry" in sys.argv
ST = S / "deliver/Supplementary_Tables/Supplementary_Tables.xlsx"
SD = S / "deliver/Source_Data/Source_Data.xlsx"
FIG = S / "deliver/Main_Figures_revised/Figure4.pdf"
OUTD = S / "deliver/Main_Figures_revised"

PAIRS = {  # pair id -> (ST19 label, 4g label, SD Material, ST3 Hap1/Hap2 gene numbers)
    "dura_final": ("TK (dura)", "TK", "dura", (33226, 33488)),
    "pisifera_final": ("NS (pisifera)", "NS", "pisifera", (33756, 33238)),
    "nrly_final": ("Nigerian", "Nigerian", "nrly", (31731, 33139)),
    "MZ4_final": ("E. oleifera", "E. oleifera", "meizhou4", (30307, 30210)),
}

res = {}
for pid, (st_lab, fig_lab, mat, st3) in PAIRS.items():
    f = SUM / f"{pid}.allele_summary.final_annot_20260923.tsv"
    row = next(csv.DictReader(open(f), delimiter="\t"))
    r = {k: int(v) for k, v in row.items() if k != "Sample"}
    assert (r["p_mRNA_gff"], r["a_mRNA_gff"]) == st3, (pid, r["p_mRNA_gff"], r["a_mRNA_gff"], st3)
    assert r["Homozygous"] + r["Heterozygous"] + r["No_hit"] == r["RBH"], pid
    res[pid] = r
    print(pid, r)
if DRY:
    sys.exit(0)

# ---- ST19 rows 33-36 + note
shutil.copy2(ST, S / "st_backup_before_G2.xlsx")
wb = openpyxl.load_workbook(ST)
ws = wb["Supplementary Table 19"]
label2row = {ws.cell(i, 1).value: i for i in range(4, ws.max_row + 1)}
for pid, (st_lab, _, _, st3) in PAIRS.items():
    i, r = label2row[st_lab], res[pid]
    rbh = r["RBH"]
    ws.cell(i, 5).value, ws.cell(i, 7).value = st3
    ws.cell(i, 8).value = rbh
    ws.cell(i, 9).value, ws.cell(i, 10).value = r["Homozygous"], round(100 * r["Homozygous"] / rbh, 2)
    ws.cell(i, 11).value, ws.cell(i, 12).value = r["Heterozygous"], round(100 * r["Heterozygous"] / rbh, 2)
    ws.cell(i, 13).value, ws.cell(i, 14).value = r["No_hit"], float(f"{100 * r['No_hit'] / rbh:.4g}")
    ws.cell(i, 15).value = round(100 * r["Homozygous"] / rbh, 2)
    assert ws.cell(i, 9).value + ws.cell(i, 11).value + ws.cell(i, 13).value == ws.cell(i, 8).value
note_row = [i for i in range(1, ws.max_row + 1) if str(ws.cell(i, 1).value or "").startswith("Note")][0]
ws.cell(note_row, 1).value = (
    "Note: FL, TN, TK, NS, Nigerian and E. oleifera use paired haplotypes (Hap1 versus Hap2); the other 27 "
    "materials use primary and alternate contigs. For TK, NS, Nigerian and E. oleifera, allele pairs were "
    "recomputed on 24 September 2026 from the final 12 August 2026 haplotype annotations with "
    "haplotype-specific gene identifiers; gene numbers are those of the final annotations (Supplementary "
    "Table 3), and gene models without a valid open reading frame were not used for pairing. TK and NS are "
    "homozygous parents; allele pairs with identical coding sequences are the majority in TK (64.9%) and a large "
    "fraction in NS (43.1%). RBH pairs include "
    "identical-CDS, divergent-CDS and other pairs. Homo_rate = identical-CDS alleles / RBH pairs × 100.")
wb.save(ST)
print("ST19 updated")

# ---- Source Data TSV (whole sheet)
wb = openpyxl.load_workbook(SD, read_only=True)
rows = list(wb["Fig.4g_alleles"].iter_rows(values_only=True))
hdr, body = list(rows[0]), [list(r) for r in rows[1:]]
mat2pid = {v[2]: k for k, v in PAIRS.items()}
new_rows = {}
for b in body:
    pid = mat2pid.get(str(b[1]))
    if pid:
        r = res[pid]
        b[3], b[4], b[5], b[6] = r["Heterozygous"], r["Unique_inter"], r["Homozygous"], r["No_hit"]
        b[7] = sum(b[3:7]); b[8] = round(100 * b[6] / b[7], 4)
        new_rows[PAIRS[pid][1]] = b
assert len(new_rows) == 4, new_rows
with open(S / "sd_fix/Fig4g_alleles.tsv", "w", newline="") as fh:
    w = csv.writer(fh, delimiter="\t"); w.writerow(hdr); w.writerows(body)
print("Fig4g_alleles.tsv written")

# ---- Fig. 4g bars (vector)
X0, X1 = 29.195, 169.325
COL = {"bi": (0.561, 0.855, 0.937), "hs": (0.902, 0.365, 0.427), "same": (0.973, 0.686, 0.631), "unres": (0.447, 0.541, 0.757)}
bak = S / "g2/Figure4_before_G2.pdf"
shutil.copy2(FIG, bak)
doc = fitz.open(bak); pg = doc[0]
words = [w for w in pg.get_text("words") if w[2] < 29.5 and 250 < w[1] < 340]
lab_y = {}
for w in words:
    lab_y.setdefault(w[4], (w[1] + w[3]) / 2)
lab_y["E. oleifera"] = next((w[1] + w[3]) / 2 for w in words if w[4] == "oleifera")
bars = {}
for dr in pg.get_drawings():
    r = dr["rect"]
    if dr.get("fill") and r.x0 < 180 and 250 < r.y0 < 340 and r.height < 3:
        bars.setdefault((round(r.y0, 2), round(r.y1, 2)), []).append(r)
targets = {}
for lab, b in new_rows.items():
    yc = lab_y[lab]
    key = min(bars, key=lambda k: abs((k[0] + k[1]) / 2 - yc))
    targets[lab] = (key, b)
for lab, (key, b) in targets.items():
    y0, y1 = key
    pg.add_redact_annot(fitz.Rect(X0 - 0.2, y0 - 0.15, X1 + 0.2, y1 + 0.15), fill=None)
pg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                    text=fitz.PDF_REDACT_TEXT_NONE)
for lab, (key, b) in targets.items():
    y0, y1 = key; tot = b[7]; x = X0
    for k, v in (("bi", b[3]), ("hs", b[4]), ("same", b[5]), ("unres", b[6])):
        w = (X1 - X0) * v / tot
        if w > 0:
            pg.draw_rect(fitz.Rect(x, y0, x + w, y1), color=None, fill=COL[k], width=0)
        x += w
doc.save(FIG, garbage=3, deflate=True)
print("Figure4.pdf saved")

# verify: every 4g row width fractions vs the updated SD rows
doc = fitz.open(FIG); pg = doc[0]
hx = lambda c: "%02x%02x%02x" % tuple(round(v * 255) for v in c)
cmap = {"8fdaef": 3, "e65d6d": 4, "f8afa1": 5, "728ac1": 6}
rowsg = {}
for dr in pg.get_drawings():
    r = dr["rect"]
    if dr.get("fill") and r.x0 < 180 and 250 < r.y0 < 340 and r.height < 3 and hx(dr["fill"]) in cmap:
        rowsg.setdefault(round((r.y0 + r.y1) / 2, 1), {})[cmap[hx(dr["fill"])]] = r.width / (X1 - X0)
labs = sorted(lab_y.items(), key=lambda t: t[1])
disp2body = {str(b[2]): b for b in body}
worst = 0
for lab, yc in labs:
    if lab not in disp2body:
        continue
    key = min(rowsg, key=lambda k: abs(k - yc)); b = disp2body[lab]
    err = max(abs(rowsg[key].get(j, 0) - b[j] / b[7]) for j in (3, 4, 5, 6))
    worst = max(worst, err)
print("max width error over all rows (fraction):", round(worst, 5))
pix = pg.get_pixmap(dpi=600, alpha=False)
Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(OUTD / "Figure4.tif", compression="tiff_lzw", dpi=(600, 600))
pg.get_pixmap(dpi=150, alpha=False).save(OUTD / "Figure4_preview.png")
pg.get_pixmap(dpi=600, alpha=False, clip=fitz.Rect(0, 235, 180, 345)).save(S / "g2/Figure4g_zoom.png")
print("exports done")
