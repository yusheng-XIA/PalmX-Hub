#!/usr/bin/env python3
"""Package BGC figure source data; run with system python3 (openpyxl available)."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT.parent / 'panBGC_visualization_revised_diversity_20260708/tables_v20260912'
TABLES = ROOT / 'tables'
inputs = []
def read(path):
    inputs.append(path)
    with path.open(newline='') as f:
        return list(csv.DictReader(f, delimiter='\t'))
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

matrix = read(TABLES / 'v5_display_matrix.tsv')
families = read(TABLES / 'v5_row_order.tsv')
materials = read(TABLES / 'v5_column_order.tsv')
summary = {r['Material']:r for r in read(SOURCE / 'st20_summary.tsv')}
raw_calls = read(SOURCE / 'st20_bgc_detail.tsv')
# The source TSV contains one exact duplicate header line; remove only that line.
calls = [r for r in raw_calls if not all(k == v for k,v in r.items())]
assert len(raw_calls)-len(calls) == 1
assert len(calls) == len({r['BGC_ID'] for r in calls}) == 814
assert len(families) == 52 and len(materials) == 33
for r in families:
    for k in ['Display_Row','Prevalence','BGC_Copies']:r[k]=int(r[k])
for i,r in enumerate(materials,1):
    r['Display_Column']=i
    r['Total_Clusters']=int(r['Total_Clusters'])
    assert r['Total_Clusters']==int(summary[r['Material']]['Total_Clusters'])
    r['Avg_Genes_Per_Cluster']=float(summary[r['Material']]['Avg_Genes_Per_Cluster'])
    r['Types']=summary[r['Material']]['Types']
for r in matrix:
    for k in r:
        if k!='PanBGC_Family':r[k]=int(r[k])
family_order={r['PanBGC_Family']:i for i,r in enumerate(families)}
material_order={r['Material']:i for i,r in enumerate(materials)}
assert [r['PanBGC_Family'] for r in matrix]==list(family_order)
assert list(matrix[0])[1:]==list(material_order)
counts=Counter((r['PanBGC_Family'],r['Variety']) for r in calls)
assert set(r['PanBGC_Family'] for r in calls)==set(family_order)
assert set(r['Variety'] for r in calls)==set(material_order)
long=[]
for row,fam in zip(matrix,families):
    f=row['PanBGC_Family']
    assert sum(row[m] for m in material_order)==fam['BGC_Copies']
    assert sum(row[m]>0 for m in material_order)==fam['Prevalence']
    for mat in material_order:
        n=row[mat]
        assert n==counts[(f,mat)]
        long.append(dict(Display_Row=fam['Display_Row'],Display_Column=material_order[mat]+1,
                         Family_ID=fam['Family_ID'],PanBGC_Family=f,Material=mat,
                         Copy_Number=n,Presence=int(n>0),Display_Level='>=2' if n>=2 else str(n)))
for mat in materials:
    m=mat['Material']
    mat['Present_Families']=sum(r[m]>0 for r in matrix)
    assert sum(r[m] for r in matrix)==mat['Total_Clusters']
assert sum(r['Copy_Number'] for r in long)==814
assert sum(r['Presence'] for r in long)==795
assert sum(r['Copy_Number']>=2 for r in long)==18
calls.sort(key=lambda r:(family_order[r['PanBGC_Family']],material_order[r['Variety']],r['BGC_ID']))
classes=[]
for c,n,tot in [('Core',9,314),('Soft-core',2,66),('Shell',29,420),('Unique',12,14)]:
    subset=[r for r in families if r['Class']==c]
    assert len(subset)==n and sum(r['BGC_Copies'] for r in subset)==tot
    classes.append(dict(Class=c,Families=n,BGC_Copies=tot,Definition_Panel='Original 34-assembly audit'))
notes=[
('Title','Source data for pan-BGC overview (v3/v4/v5; same scientific data)'),
('Contents','52 families × 33 materials; 814 BGC calls; 795 occupied cells; 18 cells with >=2 copies.'),
('Copy_number_matrix','Exact integer copy numbers 0–3, in figure display order. Values of 3 are preserved.'),
('Family_summary','Display_Row is top-to-bottom (1-based); Family_ID F# maps to stable PFAMFAM_####. Prevalence counts materials; BGC_Copies sums calls.'),
('Material_summary','Display_Column is left-to-right (1-based); Total_Clusters sums calls; Present_Families counts occupied families. Avg_Genes_Per_Cluster and Types are upstream metadata, not newly computed gene measurements.'),
('BGC_calls','814 distinct BGC calls; Variety means Material. BGC_ID uniquely identifies a call. Cluster_ID is local to a material; Canonical_Type is the source label.'),
('Matrix_long','All 1716 family–material combinations, including absences. Presence is 0/1; Display_Level collapses >=2 only for colour encoding.'),
('Class_summary','Counts summarize inherited source-panel classes, not recalculated classes for the displayed 33 materials.'),
('Class definitions','Original 34-assembly audit: core = all 34; soft-core = >=33 excluding core; unique = exactly one; shell = remaining. Current counts: 9/2/29/12.'),
('Class caveat','Core and soft-core both occur in all 33 displayed materials. F44 and F48 retain Unique labels but occur in two current materials.'),
('Type legend','Saccharide=saccharide; Cyclopeptide=cyclopeptide; Putative=putative; Fatty acid=fatty_acid or fatty_acid-polyketide; Polyketide=polyketide; Others=remaining canonical types. Counts 21/12/6/3/2/8.'),
('Type caveat','The Fatty acid legend group is not an exhaustive set of fatty-acid-containing types. Major_Type retains the uncollapsed source family annotation.'),
('Family definition','PFAM/domain-fingerprint pan-BGC families; average-linkage/Jaccard threshold 0.65 with canonical type features, per source documentation. Not BiG-SCAPE GCFs or experimentally validated functions.'),
('Order','Material order is decreasing total BGC count, with source tie order retained. Families use the supplied frozen display order. Labels are stable IDs, not consecutive display ranks.'),
('Source cleanup','Removed one exact duplicate header line from st20_bgc_detail.tsv. No BGC records dropped or deduplicated.'),
('Validation','Every matrix cell independently reconstructed from BGC_calls; all 1716 match. Family/material summaries match matrix totals.'),
('Reproduce','Run python3 scripts/build_source_data.py from the existing figure workspace. Source table filenames and SHA256 checksums are in Provenance.'),
('Status','Source-data package for review; no new clustering or biological analysis performed.')]
provenance=[dict(File=p.name,SHA256=sha(p)) for p in inputs]
sheets={
    'README':(['Field','Description'],[list(x) for x in notes]),
    'Copy_number_matrix':(list(matrix[0]),[list(r.values()) for r in matrix]),
    'Family_summary':(list(families[0]),[list(r.values()) for r in families]),
    'Material_summary':(list(materials[0]),[list(r.values()) for r in materials]),
    'BGC_calls':(list(calls[0]),[list(r.values()) for r in calls]),
    'Matrix_long':(list(long[0]),[list(r.values()) for r in long]),
    'Class_summary':(list(classes[0]),[list(r.values()) for r in classes]),
    'Provenance':(list(provenance[0]),[list(r.values()) for r in provenance])}
wb=openpyxl.Workbook();wb.remove(wb.active)
outputs=[]
for name,(header,rows) in sheets.items():
    ws=wb.create_sheet(name);ws.append(header)
    for row in rows:ws.append(row)
    ws.freeze_panes='B2' if name=='Copy_number_matrix' else 'A2'
    ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='506D69')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=30
    for col in ws.columns:
        width=min(36,max(12,max(len(str(c.value or '')) for c in col)+2))
        ws.column_dimensions[get_column_letter(col[0].column)].width=width
    if name=='README':
        ws.column_dimensions['A'].width=25;ws.column_dimensions['B'].width=105
        for row in ws.iter_rows(min_row=2):
            row[1].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=48
    if name not in ['README','Provenance']:
        p=TABLES/f'SourceData_panBGC_{name}.tsv'
        with p.open('w',newline='') as f:
            w=csv.writer(f,delimiter='\t');w.writerow(header);w.writerows(rows)
        outputs.append(p)
xlsx=ROOT/'SourceData_panBGC_overview.xlsx';wb.save(xlsx);outputs.append(xlsx)
# Reopen the workbook and verify every exported value, including integer 3s.
check=openpyxl.load_workbook(xlsx,read_only=True,data_only=True)
for name,(header,rows) in sheets.items():
    actual=list(check[name].values)
    assert actual==[tuple(header)]+[tuple(row) for row in rows],name
check.close()
report={'status':'PASS','families':52,'materials':33,'BGC_calls':814,'matrix_cells':1716,
        'occupied_cells':795,'multicopy_cells':18,'matrix_vs_calls_mismatches':0,
        'duplicate_source_header_rows_removed':1,'xlsx_readback':'all cells match',
        'sheet_rows':{k:len(v[1]) for k,v in sheets.items()},
        'input_sha256':{p.name:sha(p) for p in inputs},
        'output_sha256':{str(p.relative_to(ROOT)):sha(p) for p in outputs},
        'script_sha256':sha(Path(__file__)),'openpyxl_version':openpyxl.__version__}
(ROOT/'qc/SourceData_panBGC_QA.json').write_text(json.dumps(report,indent=2)+'\n')
(ROOT/'SourceData_panBGC_README.md').write_text('# pan-BGC source data\n\n'+ '\n\n'.join(f'**{k}**: {v}' for k,v in notes)+'\n')
print(json.dumps({'workbook':str(xlsx),'worksheets':list(sheets),'validation':'PASS: all 1716 cells match 814 BGC calls; Excel readback verified'},indent=2))
