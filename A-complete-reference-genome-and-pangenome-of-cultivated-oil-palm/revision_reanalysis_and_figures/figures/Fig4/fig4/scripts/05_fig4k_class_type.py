#!/usr/bin/env python3
"""Fig. 4k: classes recomputed on the 33 displayed materials (Core 33/33, Soft-core 32/33,
Shell 2-31/33, Unique 1/33) and F28 typed as in Source Data (transporter_associated -> Others).
Rows are re-sorted with the panel's own rule (build_panbgc_overview_v5.py::display_order:
class Unique->Core, inside a class major-type rank (first appearance F1..F52), prevalence,
copies, family order; reversed). All row-bound vector items are moved at content-stream level
(heat-map cells, prevalence dots, copy-number bars) - geometry, colours and widths unchanged;
class/type strips are re-emitted with the new colours; dashed class separators and F-row labels
are re-placed with the script's rules; legend counts updated. Column labels untouched.
in: steps/step4_4c.pdf  out: steps/step5_4k.pdf  +  sd/Fig4k_class_type.tsv
"""
import csv, re
import openpyxl, fitz
from common import *

SRC, DST = FIX / "steps/step4_4c.pdf", FIX / "steps/step5_4k.pdf"
import sys as _s
if len(_s.argv) > 2: SRC, DST = Path(_s.argv[1]), Path(_s.argv[2])
V5 = SCR / "trace/Fig4/src/bgc2/v5_row_order.tsv"
CLASS_COL = {"Core": ".569 .663 .745", "Soft-core": ".851 .745 .51", "Shell": ".58 .71 .643", "Unique": ".808 .612 .62"}
TYPE_COL = {"Saccharide": ".659 .718 .8", "Cyclopeptide": ".839 .69 .541", "Putative": ".686 .749 .608",
            "Fatty acid": ".733 .667 .796", "Polyketide": ".741 .667 .576", "Others": ".769 .788 .784"}
CELL = {".945 .949 .933": 0, ".471 .663 .608": 1, ".812 .565 .49": 2}
CLASS_ORDER = {"Core": 0, "Soft-core": 1, "Shell": 2, "Unique": 3}


def type6(t):
    return {"saccharide": "Saccharide", "cyclopeptide": "Cyclopeptide", "putative": "Putative",
            "polyketide": "Polyketide", "fatty_acid": "Fatty acid", "fatty_acid-polyketide": "Fatty acid"}.get(t, "Others")


def cls33(n):
    return "Core" if n == 33 else "Soft-core" if n == 32 else "Shell" if n >= 2 else "Unique"


def display_order(fams, meta):
    seen = []
    for f in fams:                              # source order F1..F52
        if meta[f]["type"] not in seen:
            seen.append(meta[f]["type"])
    rank = {t: i for i, t in enumerate(seen)}
    idx = {f: i for i, f in enumerate(fams)}
    o = sorted(fams, key=lambda f: (CLASS_ORDER[meta[f]["cls"]], rank[meta[f]["type"]],
                                    -meta[f]["prev"], -meta[f]["copies"], idx[f]))
    return list(reversed(o))


def label_rows(order, meta):
    blocks, start = [], 0
    for i in range(1, len(order) + 1):
        if i == len(order) or meta[order[i]]["cls"] != meta[order[start]]["cls"]:
            blocks.append((start, i - 1)); start = i
    want = sorted({r for a, b in blocks for r in (a, (a + b) // 2, b)})
    keep = []
    for r in want:
        if not any(abs(r - k) <= 3 for k in keep):
            keep.append(r)
    return blocks, keep


# ---- data -----------------------------------------------------------------------------
wb = openpyxl.load_workbook(SD_XLSX, read_only=True)
P = list(wb["Fig.4k_BGC_presence"].iter_rows(values_only=True))
S = list(wb["Fig.4k_BGC_summary"].iter_rows(values_only=True))
mats = list(P[0][1:])
pres = {r[0]: list(r[1:]) for r in P[1:]}
summ = {r[0]: dict(zip(S[0], r)) for r in S[1:]}
fams = sorted(pres, key=lambda f: int(f.split("_")[1]))
fid = lambda f: "F%d" % int(f.split("_")[1])
v5 = list(csv.DictReader(open(V5), delimiter="\t"))
old_order = [r["PanBGC_Family"] for r in v5]
old_meta = {r["PanBGC_Family"]: dict(cls=r["Class"], type=r["Major_Type"], prev=int(r["Prevalence"]),
                                     copies=int(r["BGC_Copies"])) for r in v5}
assert display_order(fams, old_meta) == old_order, "rule does not reproduce the submitted row order"
new_meta = {}
for f in fams:
    n = sum(1 for v in pres[f] if v and v > 0)
    assert n == summ[f]["Genome_count"]
    new_meta[f] = dict(cls=cls33(n), type=summ[f]["Major_Type"], prev=n, copies=int(summ[f]["BGC_count"]))
new_order = display_order(fams, new_meta)
pos = {f: i for i, f in enumerate(new_order)}

doc = fitz.open(SRC)
s = get_stream(doc)

# ---- row geometry from the class strip -------------------------------------------------------
ca = s.index("q .333 .333 .333 RG 2 J .45 w 238.543 85.791 3.8289948 157.324 re W n")
cb = s.index(" Q", s.index("re f", ca) + 0) if False else s.index("Q q .333 .333 .333 RG 2 J .45 w 256.97", ca)
cls_blk = s[ca:cb]
rows_geo = [(float(y), float(h)) for y, h in re.findall(r"238\.544 ([\d.]+) 3\.828003 (-[\d.]+) re f", cls_blk)]
assert len(rows_geo) == 52
top = [y for y, h in rows_geo]; hgt = [h for y, h in rows_geo]
cen = [y + h / 2 for y, h in rows_geo]
row_of = lambda yc: min(range(52), key=lambda r: abs(cen[r] - yc))

# old strip colours must match v5 metadata (sanity: we are moving the right rows)
cur, old_cls_cols = None, []
for m in re.finditer(r"([\d.]+ [\d.]+ [\d.]+) rg|238\.544 [\d.]+ 3\.828003 -[\d.]+ re f", cls_blk):
    if m.group(1): cur = m.group(1)
    else: old_cls_cols.append(cur)
assert old_cls_cols == [CLASS_COL[old_meta[f]["cls"]] for f in old_order]

# ---- heat-map cells ----------------------------------------------------------------------------
ha = s.index("q 261.789 85.791 198.422 157.324 re W n")
hb = s.index(" Q .2 .2 .2 rg BT", ha)
blk = s[ha:hb]
xs_label = [w for w in doc[0].get_text("words", clip=fitz.Rect(255, 536, 462, 575))]
cur, cells, out, p0 = None, {}, [], 0
tok = re.compile(r"([\d.]+ [\d.]+ [\d.]+) rg|(-?[\d.]+) (-?[\d.]+) (-?[\d.]+) (-?[\d.]+) re B")
for m in tok.finditer(blk):
    if m.group(1):
        cur = m.group(1); continue
    x, y, w_, h = (float(m.group(i)) for i in range(2, 6))
    r = row_of(y + h / 2); c = round((x - 261.79) / 6.01285)
    cells[(r, c)] = CELL[cur]
    nr = pos[old_order[r]]
    out.append(blk[p0:m.start()] + f"{m.group(2)} {num(top[nr])} {m.group(4)} {num(hgt[nr])} re B")
    p0 = m.end()
out.append(blk[p0:])
assert len(cells) == 52 * 33
s = s[:ha] + "".join(out) + s[hb:]

# check matrix against Source Data using the column labels printed under the panel
page = doc[0]
labs = []
for b in page.get_text("dict", clip=fitz.Rect(255, 536, 462, 590))["blocks"]:
    for l in b.get("lines", []):
        if l["dir"][1] < -0.5:
            labs.append((l["bbox"][0], "".join(sp["text"] for sp in l["spans"]).replace("\xa0", " ")))
labs = [t for x, t in sorted(labs)]
assert len(labs) == 33 and set(labs) == set(mats), labs
for r, f in enumerate(old_order):
    for c, m_ in enumerate(labs):
        v = pres[f][mats.index(m_)] or 0
        assert cells[(r, c)] == min(v, 2), (f, m_, cells[(r, c)], v)

# ---- class & type strips (re-emitted) ----------------------------------------------------------
def strip(x, w, colour_of):
    parts, cur = [], None
    for r, f in enumerate(new_order):
        col = colour_of(f)
        if col != cur:
            parts.append(f"{col} rg"); cur = col
        parts.append(f"{x} {num(top[r])} {w} {num(hgt[r])} re f")
    return " ".join(parts)

ca = s.index("q .333 .333 .333 RG 2 J .45 w 238.543 85.791 3.8289948 157.324 re W n")
cb = s.index(" Q q .333 .333 .333 RG 2 J .45 w 256.97", ca)
s = s[:ca] + "q .333 .333 .333 RG 2 J .45 w 238.543 85.791 3.8289948 157.324 re W n " + \
    strip("238.544", "3.828003", lambda f: CLASS_COL[new_meta[f]["cls"]]) + s[cb:]
ta = s.index("q .333 .333 .333 RG 2 J .45 w 256.97 85.791 3.402008 157.324 re W n")
tb = s.index(" Q 1 1 1 rg 2 J .45 w 465.884", ta)
old_type_blk = s[ta:tb]
cur, old_type_cols = None, []
for m in re.finditer(r"([\d.]+ [\d.]+ [\d.]+) rg|256\.97 [\d.]+ 3\.402008 -[\d.]+ re f", old_type_blk):
    if m.group(1): cur = m.group(1)
    else: old_type_cols.append(cur)
assert old_type_cols == [TYPE_COL[type6(old_meta[f]["type"])] for f in old_order]
s = s[:ta] + "q .333 .333 .333 RG 2 J .45 w 256.97 85.791 3.402008 157.324 re W n " + \
    strip("256.97", "3.402008", lambda f: TYPE_COL[type6(new_meta[f]["type"])]) + s[tb:]

# ---- prevalence dots ------------------------------------------------------------------------------
pa = s.index("q 465.883 85.791 17.289002 157.324 re W n")
pb = s.index("1 1 1 rg 2 J 0 j .5 w 491.396 243.115", pa)
blk = s[pa:pb]; out, p0, nd = [], 0, 0
for m in re.finditer(r"q 1 0 0 1 ([\d.]+) ([\d.]+) cm \.267 \.357 \.408 rg (2 J 0 j \.5 w 0 0 m \.\d+ 0 \.\d+ \.\d+ \.\d+ \.\d+ c \.\d+ \.\d+ \.\d+ (\.\d+|1\.\d+) )", blk):
    y = float(m.group(2))
    # circle drawn from its bottom point; radius from the 4th control y (~0.866)
    rad = 0.866
    r = row_of(y + rad); f = old_order[r]
    xv = float(m.group(1))
    assert abs((xv - 465.8836) / (17.289 / 33) - old_meta[f]["prev"]) < 0.05
    nr = pos[f]
    out.append(blk[p0:m.start()] + f"q 1 0 0 1 {m.group(1)} {num(y + cen[nr] - cen[r], 4)} cm .267 .357 .408 rg {m.group(3)}")
    p0 = m.end(); nd += 1
out.append(blk[p0:]); assert nd == 52
s = s[:pa] + "".join(out) + s[pb:]

# ---- copy-number bars ---------------------------------------------------------------------------
ba = s.index("q .333 .333 .333 RG 491.396 85.791 21.823975 157.324 re W n")
bb = s.index(" Q", s.index("re f", ba) + 5) if False else s.index("Q q 1 0 0 1 491.3956 85.7906 cm", ba)
blk = s[ba:bb]; out, p0, nb = [], 0, 0
for m in re.finditer(r"491\.396 ([\d.]+) ([\d.]+) (-[\d.]+) re f", blk):
    y, w_, h = float(m.group(1)), float(m.group(2)), float(m.group(3))
    r = row_of(y + h / 2); f = old_order[r]
    assert abs(w_ / (19.6645 / 50) - old_meta[f]["copies"]) < 0.1, (f, w_)
    nr = pos[f]
    out.append(blk[p0:m.start()] + f"491.396 {num(y + cen[nr] - cen[r])} {m.group(2)} {m.group(3)} re f")
    p0 = m.end(); nb += 1
out.append(blk[p0:]); assert nb == 52
s = s[:ba] + "".join(out) + s[bb:]

# ---- dashed class separators --------------------------------------------------------------------
sep = "q 1 0 0 1 256.9696 {} cm .227 .227 .227 RG 0 J 1 j .55 w 0 0 m 203.242 0 l S Q"
old_seps = re.findall(r"q 1 0 0 1 256\.9696 ([\d.]+) cm \.227 \.227 \.227 RG 0 J 1 j \.55 w 0 0 m 203\.242 0 l S Q", s)
old_blocks, _ = label_rows(old_order, old_meta)
assert [round(float(v), 2) for v in old_seps] == [round(top[b + 1], 2) for a, b in old_blocks[:-1]]
new_blocks, keep = label_rows(new_order, new_meta)
first = s.index(sep.format(old_seps[0])); last = s.index(sep.format(old_seps[-1])) + len(sep.format(old_seps[-1]))
s = s[:first] + " ".join(sep.format(num(top[b + 1], 4)) for a, b in new_blocks[:-1]) + s[last:]

# ---- row labels + legend counts (text) ---------------------------------------------------------
RS = 3.025465 / 6          # one row in text-space units (6 pt font)
old_chain = ("-1.791 -1.318 Td (F45)Tj 0 -2.521 Td (F49)Tj 0 -3.025 Td (F52)Tj 0 -7.564 Td (F35)Tj "
             "0 -7.059 Td (F15)Tj .556 -3.53 Td (F3)Tj 0 -2.017 Td (F1)Tj 35.081 27.858 Td (Prevalence)Tj")
prev_abs = (0.556 + 35.081, -(2.521 + 3.025 + 7.564 + 7.059 + 3.53 + 2.017) + 27.858)   # rel. to F45 line start
parts, px, py = [], -1.791, -1.318          # first Td is relative to the "ype" line start
lx, ly = None, None
for r in keep:
    t = fid(new_order[r]); x = 0.556 if len(t) == 2 else 0.0; y = -r * RS
    if lx is None:
        parts.append(f"{num(px + x)} {num(py + y)} Td ({t})Tj")
    else:
        parts.append(f"{num(x - lx)} {num(y - ly)} Td ({t})Tj")
    lx, ly = x, y
parts.append(f"{num(prev_abs[0] - lx)} {num(prev_abs[1] - ly)} Td (Prevalence)Tj")
s = replace_once(s, old_chain, " ".join(parts))
cc = {c: sum(1 for f in new_order if new_meta[f]["cls"] == c) for c in CLASS_COL}
tc = {t: sum(1 for f in new_order if type6(new_meta[f]["type"]) == t) for t in TYPE_COL}
assert (cc["Core"], cc["Soft-core"], cc["Shell"], cc["Unique"]) == (11, 0, 31, 10)
assert (tc["Putative"], tc["Others"]) == (5, 9) and sum(tc.values()) == 52
for old, new in ((r"(Core \(9)Tj (\))Tj", r"(Core \(%d)Tj (\))Tj" % cc["Core"]),
                 (r"(Soft-core \(2\))Tj", r"(Soft-core \(%d\))Tj" % cc["Soft-core"]),
                 (r"(Shell \(29)Tj", r"(Shell \(%d)Tj" % cc["Shell"]),
                 (r"(Unique \(12\))Tj", r"(Unique \(%d\))Tj" % cc["Unique"]),
                 (r"(Putative \(6\))Tj", r"(Putative \(%d\))Tj" % tc["Putative"]),
                 (r"(Others \(8\))Tj", r"(Others \(%d\))Tj" % tc["Others"])):
    s = replace_once(s, old, new)
for t in ("Saccharide", "Cyclopeptide", "Fatty acid", "Polyketide"):
    assert f"({t} \\({tc[t]}\\))Tj" in s, t
put_stream(doc, s)
doc.save(DST, garbage=0, deflate=True)

with open(FIX / "sd/Fig4k_class_type.tsv", "w", newline="") as fh:
    w = csv.writer(fh, delimiter="\t")
    w.writerow(["Display_row_new", "Family_ID", "PanBGC_Family", "Materials_present_of_33", "BGC_copies",
                "Class_33materials", "Class_submitted_fig", "Major_Type_SourceData", "Type_legend",
                "Type_legend_submitted_fig", "Display_row_submitted", "Row_label_shown"])
    for i, f in enumerate(new_order):
        w.writerow([i + 1, fid(f), f, new_meta[f]["prev"], new_meta[f]["copies"], new_meta[f]["cls"],
                    old_meta[f]["cls"], new_meta[f]["type"], type6(new_meta[f]["type"]),
                    type6(old_meta[f]["type"]), old_order.index(f) + 1, "yes" if i in keep else ""])
print("class counts", cc); print("type counts", tc)
print("labels", [(r + 1, fid(new_order[r])) for r in keep])
print("separators after rows", [b + 1 for a, b in new_blocks[:-1]])
print("rows moved", sum(1 for f in fams if pos[f] != old_order.index(f)))
